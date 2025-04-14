import re
from os import listdir
from pathlib import Path
from spire.doc import Document, CellCollection, DocumentObjectType, FileFormat
from bs4 import BeautifulSoup
try:
    from sqlalchemy import Select
except:
    from sqlalchemy import select as Select
from typing import Tuple
from openpyxl import Workbook, load_workbook

from database import create_session, Task


def get_num_ans_pair(cells: CellCollection, ind: int) -> Tuple[int, str]:
    num_cell = cells.get_Item(ind)
    ans_cell = cells.get_Item(ind + 1)
    number = int(num_cell.Paragraphs[0].Text[:-1])
    answer = "\n".join(ans_cell.Paragraphs[0].Text.strip().split("\x0b"))
    return (number, answer)


def parse_document_to_txt(document: Document) -> None:
    section = document.Sections[0]
    table = section.Tables[0]
    rows = table.Rows
    for i in range(rows.Count):
        row = rows.get_Item(i)
        first_cell = row.Cells[0]
        second_cell = row.Cells[1]

        task_number = first_cell.Paragraphs[0].Text

        image = None
        second_cell_par_children = second_cell.Paragraphs[0].ChildObjects
        for j in range(second_cell_par_children.Count):
            obj = second_cell_par_children.get_Item(j)
            if obj.DocumentObjectType == DocumentObjectType.Picture:
                image = obj.ImageBytes

        htm_save_path = Path("templates", f"task{task_number}.htm")
        paragraphs = second_cell.Paragraphs
        with open(htm_save_path, "w", encoding="utf-8") as htm_file:
            for i in range(paragraphs.Count):
                task_text = paragraphs.get_Item(i).Text
                htm_file.write(task_text.replace("\x0b", "\n") + "\n")

        if image is not None:
            img_save_path = Path("static", "img", f"img{task_number}.png")
            with open(img_save_path, "wb") as img_file:
                img_file.write(image)


def separate_word_to_tasks(document: Document) -> None:
    section = document.Sections[0]
    tables = section.Tables
    task_table = tables[0]
    
    task_rows = task_table.Rows
    for i in range(task_rows.Count):
        row = task_rows.get_Item(i)
        task_number = row.Cells[0].Paragraphs[0].Text

        cut_off_task = Document()
        section = cut_off_task.AddSection()
        new_table = section.AddTable()
        new_table.AddRow()
        new_table.Rows.Insert(0, row.Clone())

        path_to_img = Path("..", "static", "img")
        cut_off_task.HtmlExportOptions.ImageEmbedded = False
        cut_off_task.HtmlExportOptions.ImagesPath = str(path_to_img)

        path_to_save = Path("templates", f"task{task_number}.html")
        cut_off_task.SaveToFile(str(path_to_save), FileFormat.Html)
        cut_off_task.Close()

    try:
        answer_table = tables[1]
    except Exception as ex:
        print(ex)
        return

    answer_dict = {}
    answer_rows = answer_table.Rows
    for i in range(1, answer_rows.Count):
        row = answer_rows.get_Item(i)
        if i in range(1, 5):
            for cell_ind in range(0, 8, 2):
                number, answer = get_num_ans_pair(row.Cells, cell_ind)
                answer_dict[number] = answer
        elif i == 5 or i == 6:
            for cell_ind in range(0, 6, 2):
                if i == 5 and cell_ind == 4:
                    num_cell = row.Cells.get_Item(cell_ind)
                    ans_cell = row.Cells.get_Item(cell_ind + 1)
                    numbers = num_cell.Paragraphs[0].Text.split("\x0b")
                    answers = ans_cell.Paragraphs[0].Text.strip().split("\x0b")
                    for number, answer in zip(numbers, answers):
                        answer_dict[number[:-1]] = answer
                    continue
                number, answer = get_num_ans_pair(row.Cells, cell_ind)
                answer_dict[number] = answer
        else:
            number, answer = get_num_ans_pair(row.Cells, 0)
            answer_dict[number] = answer
    
    with create_session() as session:
        all_tasks_stm = Select(Task)
        all_tasks = session.scalars(all_tasks_stm).all()

        for task in all_tasks:
            session.delete(task)
        session.commit()
        
        try:
            for number, answer in answer_dict.items():
                new_task = Task(task_number=number, task_answer=answer)
                session.add(new_task)
            session.commit()
        except Exception as ex:
            print(ex)
            session.rollback()


def remove_red_watermarks():
    tasks = listdir("templates")
    for task in tasks:
        if not re.fullmatch(r"task(?:[1-9]|1[0-8]|2[2-7]|19 20 21)\.html", task):
            continue

        path_to_task = Path("templates", task)
        with open(path_to_task, "r", encoding="utf-8") as html_file:
            old_html = html_file.readline()

        watermark_code = """<p style="margin-top:0pt; margin-bottom:0pt; font-size:12pt"><span style="font-family:'Times New Roman'; color:#ff0000">Evaluation Warning: The document was created with Spire.Doc for Python.</span></p>"""
        old_html = old_html.replace(watermark_code, "")

        with open(path_to_task, "w", encoding="utf-8") as html_file:
            html_file.write(old_html)


def remove_task_numbers():
    tasks = listdir("templates")
    for task in tasks:
        if not re.fullmatch(r"task(?:[1-9]|1[0-8]|2[2-7]|19 20 21)\.html", task):
            continue

        path_to_task = Path("templates", task)
        with open(path_to_task, "r", encoding="utf-8") as html_file:
            old_html = html_file.readline()

        column_tags_pattern = r"<td.*?>.*?<\/td>"
        column_tags = re.findall(column_tags_pattern, old_html)
        number_tag = column_tags[0]
        old_html = old_html.replace(number_tag, "")

        with open(path_to_task, "w", encoding="utf-8") as html_file:
            html_file.write(old_html)


def split19_21():
    with open(Path("templates", "task19 20 21.html"), "rb") as tasks_html:
        soup = BeautifulSoup(tasks_html, "html.parser")
        all_spans = soup.find_all("span")

    question_indexes = {}
    for i, span in enumerate(all_spans):
        if "Вопрос 1." in span:
            question_indexes["first"] = i
        elif "Вопрос 2." in span:
            question_indexes["second"] = i
        elif "Вопрос 3." in span:
            question_indexes["third"] = i

    # need to refactor
    task_19_text = "".join(map(str, all_spans[: question_indexes["second"]])).replace(
        "Вопрос 1.", "<br>"
    )
    task_20_text = "".join(
        map(str, all_spans[question_indexes["second"] + 2 : question_indexes["third"]])
    )
    task_21_text = "".join(map(str, all_spans[question_indexes["third"] + 2 :]))

    # need to refactor
    with open(Path("templates", "task19.html"), "w", encoding="utf-8") as task19_html:
        task19_html.write(
            '<div><table cellspacing="0" cellpadding="0" style="border-collapse:collapse"><tr><td style="width:500.05pt; padding-right:5.4pt; padding-left:5.4pt; vertical-align:top"><p style="margin:4.3pt; line-height:15pt">'
            + task_19_text
            + "</p></td></tr></table></div>"
        )

    with open(Path("templates", "task20.html"), "w", encoding="utf-8") as task20_html:
        task20_html.write(
            '<div><table cellspacing="0" cellpadding="0" style="border-collapse:collapse"><tr><td style="width:500.05pt; padding-right:5.4pt; padding-left:5.4pt; vertical-align:top"><p style="margin:4.3pt; line-height:15pt">'
            + task_20_text
            + "</p></td></tr></table></div>"
        )

    with open(Path("templates", "task21.html"), "w", encoding="utf-8") as task21_html:
        task21_html.write(
            '<div><table cellspacing="0" cellpadding="0" style="border-collapse:collapse"><tr><td style="width:500.05pt; padding-right:5.4pt; padding-left:5.4pt; vertical-align:top"><p style="margin:4.3pt; line-height:15pt">'
            + task_21_text
            + "</p></td></tr></table></div>"
        )


def get_body():
    tasks = listdir("templates")
    for task in tasks:
        if not re.fullmatch(r"task(?:[1-9]|1[0-8]|2[2-7]|19 20 21)\.html", task):
            continue
        path_to_task = Path("templates", task)
        with open(path_to_task, "r", encoding="utf-8") as html_file:
            old_html = html_file.readline()

        body_pattern = r"<body>(.*?)<\/body>"
        old_html = re.findall(body_pattern, old_html)[0]

        with open(path_to_task, "w", encoding="utf-8") as html_file:
            html_file.write(old_html)


def parse_Poliacov_document(path_to_document: Path) -> None:
    document_to_parse = Document()
    document_to_parse.LoadFromFile(str(path_to_document))
    separate_word_to_tasks(document_to_parse)
    remove_red_watermarks()
    remove_task_numbers()
    get_body()
    split19_21()


def parse_from_images(path_to_images_folder: Path) -> None:
    path_to_static = Path(".", "static", "image_parser")
    path_to_templates = Path(".", "templates")
    
    images = listdir(path_to_images_folder)
    images.remove("answers.xlsx")
    for image in images:
        with open(path_to_images_folder / image, "rb") as in_image:
            with open(path_to_static / image, "wb") as out_image:
                out_image.writelines(in_image.readlines())
        with open(path_to_templates / (image[:-3] + "html"), "w", encoding="utf-8") as html_file:
            soup = BeautifulSoup("<div></div>", "html.parser")
            img_tag = soup.new_tag("img", src=f"../static/image_parser/{image}")
            soup.div.append(img_tag)
            html_file.write(str(soup))

    answer_workbook = load_workbook(path_to_images_folder / "answers.xlsx")
    active_sheet = answer_workbook.active

    with create_session() as db_session:
        try:
            for number, answer in active_sheet.rows:
                new_task = Task(task_number=number.value, task_answer=answer.value)
                db_session.add(new_task)
            db_session.commit()
        except Exception as ex:
            print(ex)
            db_session.rollback()


if __name__ == "__main__":
    # document_to_parse = Document("test.docx")
    path_to_document = Path("test2.docx")
    parse_from_images(Path("imgs_to_parse"))
